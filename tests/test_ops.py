"""One test (or more) per operation. All inputs come from fixtures in conftest.py."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import pypdf
import pytest
from openpyxl import load_workbook

import pdfops
from pdfops.table_ops import to_number


def ok(r: dict) -> dict:
    """Assert the report envelope, JSON-serializability and success."""
    assert set(r) == {"ok", "changed", "warnings", "details"}, r
    json.dumps(r)
    assert r["ok"], r["details"]
    return r


def failed(r: dict, error_type: str | None = None) -> dict:
    assert set(r) == {"ok", "changed", "warnings", "details"}, r
    json.dumps(r)
    assert r["ok"] is False and r["changed"] is False, r
    assert r["details"]["error"]
    if error_type:
        assert r["details"]["error_type"] == error_type, r["details"]
    return r


def page_texts(path) -> list[str]:
    return [p.extract_text() or "" for p in pypdf.PdfReader(str(path)).pages]


# ---------------------------------------------------------------- framework


def test_schemas_registered_valid_and_short():
    import jsonschema
    assert len(pdfops.SCHEMAS) == 21
    assert set(pdfops.SCHEMAS) == set(pdfops.OPERATIONS)
    for name, s in pdfops.SCHEMAS.items():
        assert len(json.dumps(s, separators=(",", ":"))) <= 600, name
        jsonschema.Draft202012Validator.check_schema(s)


def test_never_raises_on_bad_input(tmp_path):
    failed(pdfops.info(path=str(tmp_path / "missing.pdf")), "not_found")
    failed(pdfops.run("info", {"path": 5}), "invalid_input")
    failed(pdfops.run("nope", {}), "unknown_operation")
    failed(pdfops.extract_pages("x.pdf"), "invalid_input")
    junk = tmp_path / "junk.pdf"
    junk.write_bytes(b"this is not a pdf")
    failed(pdfops.info(path=str(junk)), "not_pdf")
    trunc = tmp_path / "trunc.pdf"
    trunc.write_bytes(b"%PDF-1.4\n1 0 obj << /Type /Catalog")
    assert pdfops.info(path=str(trunc))["ok"] is False


def test_no_network_or_subprocess_in_library():
    banned = ("subprocess", "socket", "urllib", "http.client", "requests", "os.system", "Popen")
    for f in Path(pdfops.__file__).parent.glob("*.py"):
        src = f.read_text(encoding="utf-8")
        for word in banned:
            assert word not in src, f"{word} in {f.name}"


# ---------------------------------------------------------------- inspect


def test_info(text_pdf):
    d = ok(pdfops.info(path=str(text_pdf)))["details"]
    assert d["pages"] == 3 and d["encrypted"] is False and d["has_form"] is False
    assert d["page_details"][0]["name"] == "Letter"
    assert d["page_details"][0]["kind"] == "text"
    assert d["metadata"]["Title"] == "Test Report"
    assert d["document_kind"] == "text"


def test_info_encrypted_reported_cleanly(encrypted_pdf):
    r = ok(pdfops.info(path=str(encrypted_pdf)))
    assert r["details"]["encrypted"] is True and r["details"]["opened"] is False
    assert "password" in r["details"]["open_error"]
    assert any("encrypted" in w for w in r["warnings"])
    # other operations fail cleanly without the password ...
    failed(pdfops.extract_text(path=str(encrypted_pdf)), "encrypted")
    failed(pdfops.extract_text(path=str(encrypted_pdf), password="wrong"), "encrypted")
    # ... and work with it
    d = ok(pdfops.info(path=str(encrypted_pdf), password="sesame"))["details"]
    assert d["opened"] and d["pages"] == 1
    t = ok(pdfops.extract_text(path=str(encrypted_pdf), password="sesame"))
    assert "Confidential crop plan" in t["details"]["pages"][0]["text"]


def test_owner_only_encryption_opens_without_password(tmp_path, out):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_encryption(owner_password="owner-only")  # empty user password
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.cell(0, 10, "Restricted but readable")
    src = tmp_path / "owner.pdf"
    pdf.output(str(src))
    d = ok(pdfops.info(path=str(src)))["details"]
    assert d["encrypted"] and d["opened"] and d["pages"] == 1
    r = ok(pdfops.extract_pages(path=str(src), pages=[1], output=out("dec.pdf")))
    assert any("WITHOUT encryption" in w for w in r["warnings"])
    assert not pypdf.PdfReader(out("dec.pdf")).is_encrypted


def test_extract_text_plain_and_layout(text_pdf):
    d = ok(pdfops.extract_text(path=str(text_pdf), pages="2-3"))["details"]
    assert [p["page"] for p in d["pages"]] == [2, 3]
    assert "Drip lines" in d["pages"][0]["text"]
    lay = ok(pdfops.extract_text(path=str(text_pdf), pages=[1], mode="layout"))["details"]
    assert "Quarterly Report" in lay["pages"][0]["text"]
    short = ok(pdfops.extract_text(path=str(text_pdf), max_chars=10))
    assert short["details"]["truncated"] and sum(p["chars"] for p in short["details"]["pages"]) == 10
    failed(pdfops.extract_text(path=str(text_pdf), pages="9"), "bad_pages")


def test_extract_text_scanned_is_honest(scanned_pdf):
    r = ok(pdfops.extract_text(path=str(scanned_pdf)))
    pages = r["details"]["pages"]
    assert pages[1]["kind"] == "scanned" and pages[1]["text"] == ""
    assert "OCR" in pages[1]["note"]
    assert any("does not perform OCR" in w for w in r["warnings"])


def test_search_terms_and_regex(text_pdf):
    d = ok(pdfops.search(path=str(text_pdf), terms=["jane doe", "drip"]))["details"]
    assert d["counts"] == {"jane doe": 2, "drip": 1}
    assert d["pages_with_matches"] == [1, 2, 3]
    assert "Jane Doe" in d["matches"][0]["snippet"]
    rx = ok(pdfops.search(path=str(text_pdf), terms=r"INV-\d{4}", regex=True))["details"]
    assert [m["match"] for m in rx["matches"]] == ["INV-2041", "INV-2077"]
    assert [m["page"] for m in rx["matches"]] == [1, 3]
    miss = ok(pdfops.search(path=str(text_pdf), terms="zucchini"))
    assert miss["details"]["total_matches"] == 0 and miss["warnings"]
    failed(pdfops.search(path=str(text_pdf), terms="(", regex=True), "bad_regex")


def test_scanned_pages(scanned_pdf):
    r = ok(pdfops.scanned_pages(path=str(scanned_pdf)))
    assert r["details"]["scanned_pages"] == [2, 3]
    assert r["details"]["per_page"][0]["kind"] == "text"
    assert r["details"]["per_page"][3]["kind"] == "empty"
    assert any("OCR" in w for w in r["warnings"])
    i = ok(pdfops.info(path=str(scanned_pdf)))["details"]
    assert i["document_kind"] == "mixed" and i["scanned_pages"] == [2, 3]


# ---------------------------------------------------------------- pages


def test_extract_pages(text_pdf, out):
    r = ok(pdfops.extract_pages(path=str(text_pdf), pages="3,1", output=out("x.pdf")))
    assert r["changed"] and r["details"]["output_pages"] == 2
    t = page_texts(out("x.pdf"))
    assert "Third page" in t[0] and "Quarterly" in t[1]
    # refuses to overwrite without permission
    failed(pdfops.extract_pages(path=str(text_pdf), pages=[1], output=out("x.pdf")), "output_exists")
    ok(pdfops.extract_pages(path=str(text_pdf), pages=[1], output=out("x.pdf"), overwrite=True))
    # in-place edit is allowed only with overwrite=true
    failed(pdfops.rotate(path=out("x.pdf"), degrees=180, output=out("x.pdf")), "output_exists")
    ok(pdfops.rotate(path=out("x.pdf"), degrees=180, output=out("x.pdf"), overwrite=True))
    assert pypdf.PdfReader(out("x.pdf")).pages[0].rotation == 180


def test_delete_pages(text_pdf, out):
    r = ok(pdfops.delete_pages(path=str(text_pdf), pages=[2], output=out("d.pdf")))
    assert r["details"]["output_pages"] == 2 and r["details"]["kept"] == "1,3"
    assert "Third page" in page_texts(out("d.pdf"))[1]
    failed(pdfops.delete_pages(path=str(text_pdf), pages="1-last", output=out("e.pdf")), "bad_pages")
    assert not Path(out("e.pdf")).exists()


def test_split_modes(make_pdf, tmp_path):
    src = make_pdf("doc", 5)
    d1 = tmp_path / "r"; d1.mkdir()
    r = ok(pdfops.split(path=str(src), output_dir=str(d1), ranges=["1-2", "4-5"]))
    assert [f["page_count"] for f in r["details"]["files"]] == [2, 2]
    assert any("not in any part" in w for w in r["warnings"])
    d2 = tmp_path / "e"; d2.mkdir()
    r = ok(pdfops.split(path=str(src), output_dir=str(d2), every=2))
    assert [f["pages"] for f in r["details"]["files"]] == ["1-2", "3-4", "5"]
    d3 = tmp_path / "h"; d3.mkdir()
    r = ok(pdfops.split(path=str(src), output_dir=str(d3), halves=True))
    files = r["details"]["files"]
    assert [f["page_count"] for f in files] == [3, 2]
    assert "Page 4 of doc" in page_texts(files[1]["path"])[0]
    failed(pdfops.split(path=str(src), output_dir=str(d3), every=2, halves=True))


def test_merge_with_bookmarks(make_pdf, out):
    a, b = make_pdf("alpha", 2), make_pdf("beta", 3)
    r = ok(pdfops.merge(inputs=[str(a), {"path": str(b), "pages": "2-3", "bookmark": "Beta part"}],
                        output=out("m.pdf"), bookmarks=True))
    assert r["details"]["output_pages"] == 4
    assert r["details"]["sources"][1]["starts_at_page"] == 3
    rd = pypdf.PdfReader(out("m.pdf"))
    titles = [(o.title, rd.get_destination_page_number(o)) for o in rd.outline]
    assert titles == [("alpha", 0), ("Beta part", 2)]
    assert "Page 2 of beta" in rd.pages[2].extract_text()


def test_append_pages(make_pdf, out):
    a, b = make_pdf("main", 3), make_pdf("extra", 2)
    r = ok(pdfops.append_pages(path=str(a), source=str(b), output=out("ap.pdf"), at=2, pages=[2]))
    assert r["details"]["inserted_at_page"] == 2
    t = page_texts(out("ap.pdf"))
    assert len(t) == 4 and "Page 2 of extra" in t[1] and "Page 2 of main" in t[2]
    ok(pdfops.append_pages(path=str(a), source=str(b), output=out("ap2.pdf")))
    assert "Page 2 of extra" in page_texts(out("ap2.pdf"))[-1]


def test_insert_appendix(make_pdf, out):
    a = make_pdf("body", 2, size=(595.28, 841.89))  # A4
    b = make_pdf("annex", 2)
    r = ok(pdfops.insert_appendix(path=str(a), appendix=str(b), output=out("ia.pdf"),
                                  title="Appendix A: Soil tests"))
    assert r["details"]["appendix_starts_at_page"] == 3
    rd = pypdf.PdfReader(out("ia.pdf"))
    assert len(rd.pages) == 5
    sep = rd.pages[2]
    assert "Appendix A" in sep.extract_text()
    assert round(float(sep.mediabox.width)) == 595 and round(float(sep.mediabox.height)) == 842
    assert rd.outline[0].title == "Appendix A: Soil tests"
    assert rd.get_destination_page_number(rd.outline[0]) == 2


def test_rotate(text_pdf, out):
    r = ok(pdfops.rotate(path=str(text_pdf), degrees=90, pages="1,3", output=out("r.pdf")))
    assert [x["rotation"] for x in r["details"]["rotated"]] == [90, 90]
    rot = [p.rotation for p in pypdf.PdfReader(out("r.pdf")).pages]
    assert rot == [90, 0, 90]
    assert ok(pdfops.info(path=out("r.pdf")))["details"]["page_details"][0]["width_pt"] == 792
    failed(pdfops.rotate(path=str(text_pdf), degrees=45, output=out("r2.pdf")), "invalid_input")


def test_drop_blank_pages(scanned_pdf, out):
    dry = ok(pdfops.drop_blank_pages(path=str(scanned_pdf), dry_run=True))
    assert dry["details"]["blank_pages"] == [3, 4] and dry["changed"] is False
    r = ok(pdfops.drop_blank_pages(path=str(scanned_pdf), output=out("nb.pdf")))
    assert r["changed"] and r["details"]["output_pages"] == 2
    kinds = [p["kind"] for p in ok(pdfops.info(path=out("nb.pdf")))["details"]["page_details"]]
    assert kinds == ["text", "scanned"]


# ---------------------------------------------------------------- forms


def test_form_fields(form_pdf):
    d = ok(pdfops.form_fields(path=str(form_pdf)))["details"]
    f = {x["name"]: x for x in d["fields"]}
    assert f["grower_name"]["type"] == "text"
    assert f["organic"]["type"] == "checkbox" and f["organic"]["options"] == ["Yes"]
    assert f["organic"]["checked"] is False
    assert f["size"]["type"] == "radio" and f["size"]["options"] == ["Small", "Large"]
    assert f["crop"]["type"] == "combo" and f["crop"]["options"] == ["Kale", "Chard", "Beets"]
    assert all(x["pages"] == [1] for x in d["fields"])
    assert ok(pdfops.info(path=str(form_pdf)))["details"]["has_form"] is True


def test_fill_form_checkbox_radio_unknown(form_pdf, out):
    r = ok(pdfops.fill_form(path=str(form_pdf), output=out("f.pdf"), values={
        "grower_name": "Olive Farms", "organic": True, "size": "large", "crop": "Chard",
        "growr_name": "typo"}))
    assert r["details"]["need_appearances"] is True
    assert r["details"]["unknown_fields"][0]["name"] == "growr_name"
    assert "grower_name" in r["details"]["unknown_fields"][0]["did_you_mean"]
    assert any("unknown field" in w for w in r["warnings"])
    rd = pypdf.PdfReader(out("f.pdf"))
    assert bool(rd.trailer["/Root"]["/AcroForm"]["/NeedAppearances"]) is True
    f = {x["name"]: x for x in ok(pdfops.form_fields(path=out("f.pdf")))["details"]["fields"]}
    assert f["grower_name"]["value"] == "Olive Farms"
    assert f["organic"]["value"] == "Yes" and f["organic"]["checked"]
    assert f["size"]["value"] == "Large"
    assert f["crop"]["value"] == "Chard"
    annots = [a.get_object() for a in rd.pages[0]["/Annots"]]
    states = {str(a.get("/T", a.get("/Parent", {}).get("/T"))) + ":" + str(a.get("/Rect")[0]): a.get("/AS")
              for a in annots if "/AS" in a}
    assert "/Yes" in states.values() and "/Large" in states.values() and "/Small" not in states.values()
    # unchecking works, bad radio value is reported, all-unknown fails cleanly
    r2 = ok(pdfops.fill_form(path=out("f.pdf"), output=out("g.pdf"),
                             values={"organic": "off", "size": "Medium"}))
    assert r2["details"]["skipped"] == ["size"]
    assert any("not a radio option" in w for w in r2["warnings"])
    f2 = {x["name"]: x for x in ok(pdfops.form_fields(path=out("g.pdf")))["details"]["fields"]}
    assert f2["organic"]["checked"] is False and f2["size"]["value"] == "Large"
    failed(pdfops.fill_form(path=str(form_pdf), output=out("h.pdf"), values={"nope": 1}),
           "nothing_filled")


def test_flatten_form(form_pdf, out):
    ok(pdfops.fill_form(path=str(form_pdf), output=out("f.pdf"),
                        values={"grower_name": "Olive Farms", "organic": True}))
    r = ok(pdfops.flatten_form(path=out("f.pdf"), output=out("flat.pdf")))
    assert r["details"]["widgets_flattened"] >= 4
    rd = pypdf.PdfReader(out("flat.pdf"))
    assert "/AcroForm" not in rd.trailer["/Root"]
    assert not rd.pages[0].get("/Annots")
    assert "Olive Farms" in rd.pages[0].extract_text()
    assert ok(pdfops.form_fields(path=out("flat.pdf")))["details"]["count"] == 0
    failed(pdfops.flatten_form(path=out("flat.pdf"), output=out("again.pdf")), "no_form")


# ---------------------------------------------------------------- tables


def test_tables_two_page_merge(table_pdf):
    assert ok(pdfops.info(path=str(table_pdf)))["details"]["pages"] == 2
    d = ok(pdfops.tables(path=str(table_pdf)))["details"]
    assert d["count"] == 1
    t = d["tables"][0]
    assert t["pages"] == [1, 2] and t["continued_across_pages"]
    assert t["header"] == ["Crop", "Bed", "Pounds", "Price"]
    assert t["n_rows"] == 45  # repeated header on page 2 was dropped
    assert t["rows"][0] == ["Crop 1", "B01", "12.5", "$1.25"]
    assert t["rows"][-1][0] == "Crop 45"
    unmerged = ok(pdfops.tables(path=str(table_pdf), merge_continued=False))["details"]
    assert unmerged["count"] == 2


def test_tables_to_csv(table_pdf, tmp_path):
    d = ok(pdfops.tables_to_csv(path=str(table_pdf), output_dir=str(tmp_path)))["details"]
    assert d["count"] == 1
    with open(d["files"][0]["path"], encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    assert rows[0] == ["Crop", "Bed", "Pounds", "Price"] and len(rows) == 46


def test_tables_to_xlsx(table_pdf, out):
    d = ok(pdfops.tables_to_xlsx(path=str(table_pdf), output=out("t.xlsx")))["details"]
    assert d["sheets"][0]["rows"] == 45
    ws = load_workbook(out("t.xlsx")).worksheets[0]
    assert [c.value for c in ws[1]] == ["Crop", "Bed", "Pounds", "Price"]
    assert ws["C2"].value == 12.5 and isinstance(ws["C2"].value, float)
    assert ws["D2"].value == 1.25
    assert ws["C9"].value == 100 and ws["B2"].value == "B01"
    assert ws["C46"].value == 562.5


def test_to_number():
    assert to_number("1,234.5") == (1234.5, False)
    assert to_number("$12") == (12, False)
    assert to_number("(3.5)") == (-3.5, False)
    assert to_number("12%") == (0.12, True)
    assert to_number("-4") == (-4, False)
    assert to_number("00123") == (None, False)
    assert to_number("12a") == (None, False)
    assert to_number("1,23") == (None, False)


# ---------------------------------------------------------------- generated content


def test_summary_note(make_pdf, out):
    src = make_pdf("land", 2, size=(792, 612))  # landscape letter
    r = ok(pdfops.summary_note(path=str(src), output=out("s.pdf"), title="Summary",
                               text="Two pages.\n\nSecond paragraph with “quotes” — dash."))
    assert r["details"]["pages_added"] == 1
    rd = pypdf.PdfReader(out("s.pdf"))
    assert len(rd.pages) == 3
    assert (float(rd.pages[0].mediabox.width), float(rd.pages[0].mediabox.height)) == (792, 612)
    assert "Summary" in rd.pages[0].extract_text()
    ok(pdfops.summary_note(path=str(src), output=out("e.pdf"), text="End note", position="end"))
    assert "End note" in pypdf.PdfReader(out("e.pdf")).pages[-1].extract_text()


def test_replace_name_overlay(text_pdf, scanned_pdf, out):
    r = ok(pdfops.replace_name_overlay(path=str(text_pdf), find="Jane Doe", replace="Ann Lee",
                                       output=out("o.pdf")))
    d = r["details"]
    assert d["count"] == 2 and [x["page"] for x in d["replaced"]] == [1, 3]
    assert all(x["font_size"] == 12 for x in d["replaced"])
    assert d["pages_without_match"] == [2]
    assert any("NOT a redaction" in w for w in r["warnings"])
    text = page_texts(out("o.pdf"))[0]
    assert "Ann Lee" in text
    # honest: the old text is still in the text layer
    assert "Jane Doe" in text
    miss = failed(pdfops.replace_name_overlay(path=str(scanned_pdf), find="Jane", replace="X",
                                              output=out("n.pdf")), "not_found_in_document")
    reasons = {u["page"]: u["reason"] for u in miss["details"]["not_replaced"]}
    assert "no text layer" in reasons[2] and "OCR" in reasons[2]


def test_letter(out):
    r = ok(pdfops.letter(output=out("l.pdf"), fields={
        "sender_name": "Olive Stewart", "sender_address": "123 Grove Rd\nFallbrook, CA 92028",
        "recipient_name": "Jana Farms", "recipient_address": ["PO Box 1", "Fallbrook, CA"],
        "date": "October 27, 2026", "subject": "Seed order",
        "body": "Please confirm the order.\n\nThank you."}))
    assert r["details"]["pages"] == 1
    rd = pypdf.PdfReader(out("l.pdf"))
    assert len(rd.pages) == 1
    t = rd.pages[0].extract_text()
    for s in ("Olive Stewart", "Dear Jana Farms,", "Re: Seed order", "Sincerely,", "Thank you."):
        assert s in t
    long = ok(pdfops.letter(output=out("long.pdf"), fields={"body": "Lorem ipsum dolor sit. " * 300}))
    assert long["details"]["font_size"] < 11 and long["warnings"]
    assert len(pypdf.PdfReader(out("long.pdf")).pages) == 1
    failed(pdfops.letter(output=out("x.pdf"), fields={"body": "word " * 5000}), "too_long")
    failed(pdfops.letter(output=out("y.pdf"), fields={"body": "hi", "bogus": "x"}))
    failed(pdfops.letter(output=out("z.pdf"), fields={"subject": "no body"}), "invalid_input")
