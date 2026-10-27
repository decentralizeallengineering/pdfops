"""Table extraction (pdfplumber) with header detection and cross-page merging."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from ._analysis import NO_OCR
from ._core import (BOOL, PAGES, STR, Doc, Result, UserError, _check_target, atomic_write,
                    operation, parse_pages, ranges_text, schema)

_SETTINGS = {
    "lines": {"vertical_strategy": "lines", "horizontal_strategy": "lines"},
    "text": {"vertical_strategy": "text", "horizontal_strategy": "text"},
}
_NUM_RE = re.compile(r"^[+-]?(\d{1,3}(,\d{3})+|\d+)?(\.\d+)?$")
EDGE_BAND = 0.08  # top/bottom fraction treated as running header/footer


def to_number(s: str) -> tuple[int | float | None, bool]:
    """Parse '1,234.5', '$12', '(3.0)', '-4', '12%'. Returns (value, is_percent).

    Codes with leading zeros ('00123') stay text.
    """
    t = s.strip()
    if not t:
        return None, False
    neg = False
    if t.startswith("(") and t.endswith(")"):
        neg, t = True, t[1:-1].strip()
    pct = t.endswith("%")
    if pct:
        t = t[:-1].strip()
    if t[:1] in "+-":
        neg ^= t[0] == "-"
        t = t[1:].strip()
    for sym in ("$", "€", "£", "¥"):
        if t.startswith(sym):
            t = t[len(sym):].strip()
            break
    if t.startswith("-"):
        neg, t = not neg, t[1:]
    if not t or not _NUM_RE.match(t) or t == ".":
        return None, False
    if re.match(r"^0\d", t.replace(",", "")):
        return None, False
    t = t.replace(",", "")
    val: int | float = float(t) if "." in t else int(t)
    if neg:
        val = -val
    if pct:
        val = val / 100
    return val, pct


def _clean(cell) -> str:
    if cell is None:
        return ""
    return re.sub(r"\s*\n\s*", " ", str(cell)).strip()


def _bold_fraction(page, bbox) -> float | None:
    try:
        chars = page.within_bbox(bbox).chars
    except Exception:
        return None
    chars = [c for c in chars if c.get("text", "").strip()]
    if not chars:
        return None
    return sum("bold" in str(c.get("fontname", "")).lower() for c in chars) / len(chars)


def _header_guess(rows: list[list[str]], page=None, table=None) -> bool:
    if len(rows) < 2:
        return False
    first = rows[0]
    filled = [c for c in first if c]
    if len(filled) < max(1, (len(first) + 1) // 2):
        return False
    if any(to_number(c)[0] is not None for c in filled):
        return False
    if page is not None and table is not None and len(table.rows) >= 2:
        b0 = _bold_fraction(page, table.rows[0].bbox)
        b1 = _bold_fraction(page, table.rows[1].bbox)
        if b0 is not None and b1 is not None and b0 >= 0.5 > b1:
            return True
    body_numeric = any(to_number(c)[0] is not None for r in rows[1:] for c in r)
    if body_numeric:
        return True
    return len(set(filled)) == len(filled) and len(filled) == len(first)


def _text_between(page, top: float, bottom: float, x0: float, x1: float) -> bool:
    if bottom - top < 1:
        return False
    return any(c.get("text", "").strip() and c["top"] >= top - 0.5 and c["bottom"] <= bottom + 0.5
               and c["x1"] > x0 and c["x0"] < x1 for c in page.chars)


def extract_tables(doc: Doc, pages, strategy: str, header, merge: bool, res: Result) -> list[dict]:
    idx = parse_pages(pages, doc.n, allow_dupes=False)
    raw = []
    scanned = []
    with doc.plumber() as pdf:
        for i in idx:
            pg = pdf.pages[i]
            found = pg.find_tables(_SETTINGS[strategy])
            if not found and not pg.chars and pg.images:
                scanned.append(i + 1)
            for j, t in enumerate(found):
                rows = [[_clean(c) for c in r] for r in t.extract()]
                rows = [r for r in rows if any(r)]
                if not rows or (len(rows) == 1 and len(rows[0]) == 1):
                    continue
                ncols = max(len(r) for r in rows)
                rows = [r + [""] * (ncols - len(r)) for r in rows]
                h = _header_guess(rows, pg, t) if header == "auto" else bool(header)
                x0, top, x1, bottom = (float(v) for v in t.bbox)
                pw, ph = float(pg.width), float(pg.height)
                raw.append({
                    "page": i + 1, "pos": j, "last_on_page": j == len(found) - 1,
                    "ncols": ncols, "bbox": [round(x0, 1), round(top, 1), round(x1, 1), round(bottom, 1)],
                    "header": rows[0] if h else None, "rows": rows[1:] if h else rows,
                    "header_detected": h,
                    "gap_below": _text_between(pg, bottom, ph * (1 - EDGE_BAND), x0, x1),
                    "gap_above": _text_between(pg, ph * EDGE_BAND, top, x0, x1),
                    "pw": pw,
                })
    if scanned:
        res.warn(f"page(s) {', '.join(map(str, scanned))} are image-only; tables there cannot be "
                 f"read ({NO_OCR})")

    tables: list[dict] = []
    for t in raw:
        prev = tables[-1] if tables else None
        if (merge and prev is not None and prev["_last"]["last_on_page"]
                and t["pos"] == 0 and t["page"] == prev["pages"][-1] + 1
                and t["ncols"] == prev["n_cols"]
                and abs(t["bbox"][0] - prev["_last"]["bbox"][0]) <= 0.03 * t["pw"]
                and abs(t["bbox"][2] - prev["_last"]["bbox"][2]) <= 0.03 * t["pw"]
                and not prev["_last"]["gap_below"] and not t["gap_above"]):
            cont_rows = t["rows"]
            if t["header"] is not None:
                if prev["header"] is not None and t["header"] == prev["header"]:
                    pass  # repeated header row dropped
                elif prev["header"] is None:
                    cont_rows = [t["header"]] + cont_rows
                else:
                    cont_rows = None  # different header: a new table
            elif prev["header"] is not None and t["rows"] and t["rows"][0] == prev["header"]:
                cont_rows = t["rows"][1:]
            if cont_rows is not None:
                prev["rows"].extend(cont_rows)
                prev["pages"].append(t["page"])
                prev["_last"] = t
                continue
        tables.append({"pages": [t["page"]], "header": t["header"], "rows": list(t["rows"]),
                       "n_cols": t["ncols"], "bbox_first_page": t["bbox"], "_last": t})
    for k, t in enumerate(tables, start=1):
        t.pop("_last")
        t["id"] = k
        t["n_rows"] = len(t["rows"])
        t["continued_across_pages"] = len(t["pages"]) > 1
    return tables


_TABLE_PROPS = dict(pages=PAGES, strategy={"enum": ["lines", "text"]},
                    header={"enum": ["auto", True, False]}, merge_continued=BOOL, password=STR)


@operation(schema("Extract tables (ruled lines by default; strategy 'text' for unruled). Detects headers; merges tables continuing onto the next page.",
                  ["path"], path=STR, **_TABLE_PROPS))
def tables(res: Result, path: str, pages=None, strategy: str = "lines", header="auto",
           merge_continued: bool = True, password: str | None = None) -> None:
    doc = Doc(path, password)
    found = extract_tables(doc, pages, strategy, header, merge_continued, res)
    if not found:
        res.warn("no tables found" + ("; try strategy='text' for tables without ruling lines"
                                      if strategy == "lines" else ""))
    res.details.update({"count": len(found), "tables": found, "strategy": strategy})


def _table_label(t: dict) -> str:
    return ranges_text([p - 1 for p in t["pages"]]).replace(",", "_")


@operation(schema("Write each extracted table to its own UTF-8 CSV (header row first) in output_dir.",
                  ["path", "output_dir"], path=STR, output_dir=STR, prefix=STR, overwrite=BOOL,
                  **_TABLE_PROPS))
def tables_to_csv(res: Result, path: str, output_dir: str, prefix: str | None = None,
                  overwrite: bool = False, pages=None, strategy: str = "lines", header="auto",
                  merge_continued: bool = True, password: str | None = None) -> None:
    odir = Path(output_dir).expanduser()
    if not odir.is_dir():
        raise UserError(f"output_dir does not exist: {output_dir}", "not_found")
    doc = Doc(path, password)
    found = extract_tables(doc, pages, strategy, header, merge_continued, res)
    if not found:
        raise UserError("no tables found; nothing written", "no_tables")
    stem = prefix or Path(path).stem
    plan = []
    for t in found:
        target = odir / f"{stem}_table{t['id']:02d}_p{_table_label(t)}.csv"
        _check_target(target, overwrite)
        plan.append((target, t))
    files = []
    for target, t in plan:
        buf = io.StringIO(newline="")
        wr = csv.writer(buf, lineterminator="\r\n")
        if t["header"]:
            wr.writerow(t["header"])
        wr.writerows(t["rows"])
        atomic_write(target, buf.getvalue().encode("utf-8-sig"))
        files.append({"path": str(target), "table": t["id"], "pages": t["pages"],
                      "rows": t["n_rows"], "cols": t["n_cols"], "has_header": t["header"] is not None})
    res.changed = True
    res.details.update({"files": files, "count": len(files)})


@operation(schema("Write tables to one .xlsx, one sheet per table; numbers stored as numbers.",
                  ["path", "output"], path=STR, output=STR, overwrite=BOOL, **_TABLE_PROPS))
def tables_to_xlsx(res: Result, path: str, output: str, overwrite: bool = False, pages=None,
                   strategy: str = "lines", header="auto", merge_continued: bool = True,
                   password: str | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    out = Path(output).expanduser()
    if out.suffix.lower() != ".xlsx":
        raise UserError("output must end in .xlsx")
    _check_target(out, overwrite)
    doc = Doc(path, password)
    found = extract_tables(doc, pages, strategy, header, merge_continued, res)
    if not found:
        raise UserError("no tables found; nothing written", "no_tables")
    wb = Workbook()
    wb.remove(wb.active)
    sheets = []
    for t in found:
        title = f"Table{t['id']} p{_table_label(t)}"[:31]
        ws = wb.create_sheet(title=re.sub(r"[\[\]:*?/\\]", "_", title))
        widths = [0] * t["n_cols"]
        r0 = 1
        if t["header"]:
            for c, v in enumerate(t["header"], start=1):
                ws.cell(row=1, column=c, value=v).font = Font(bold=True)
                widths[c - 1] = max(widths[c - 1], len(v))
            ws.freeze_panes = "A2"
            r0 = 2
        numeric = text = 0
        for r, row in enumerate(t["rows"], start=r0):
            for c, v in enumerate(row, start=1):
                num, pct = to_number(v)
                cell = ws.cell(row=r, column=c, value=num if num is not None else (v or None))
                if num is not None:
                    numeric += 1
                    if pct:
                        cell.number_format = "0.00%"
                elif v:
                    text += 1
                widths[c - 1] = max(widths[c - 1], len(v))
        for c, wdt in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = min(60, max(8, wdt + 2))
        sheets.append({"sheet": ws.title, "table": t["id"], "pages": t["pages"],
                       "rows": t["n_rows"], "cols": t["n_cols"], "numeric_cells": numeric,
                       "text_cells": text})
    buf = io.BytesIO()
    wb.save(buf)
    atomic_write(out, buf.getvalue())
    res.changed = True
    res.details.update({"output": str(out), "sheets": sheets, "count": len(sheets)})
